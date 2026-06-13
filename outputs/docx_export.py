"""Word export (outputs/docx_export.py). Includes audit trail per finding."""

import io


def to_docx(findings: list, source) -> bytes:
    """Render findings (with audit trails) to a Word .docx byte string."""
    from docx import Document
    doc = Document()
    doc.add_heading("TRACE-AI Findings Report", 0)
    src_label = ", ".join(source) if isinstance(source, list) else source
    doc.add_paragraph(f"Source document(s): {src_label}")
    doc.add_paragraph("Golden Rule: No Evidence = No Claim")
    for f in findings:
        c = f["confidence"]
        pct = c.get("percent", round(c.get("score", 0) * 100))
        doc.add_heading(f["theme"], level=1)
        doc.add_paragraph(f"Status: {f['status']}")
        conf_line = f"Confidence: {c['band']} ({pct}%)"
        if c.get("partial_strength"):
            conf_line += f" · {c['partial_strength']}"
        doc.add_paragraph(conf_line)
        if f.get("reason"):
            doc.add_paragraph(f"Reason: {f['reason']}")
        for ev in f.get("evidence", []):
            src = ev.get("source_document", "")
            etype = ev.get("evidence_type", {}).get("label", "")
            doc.add_paragraph(
                f"[{etype}] \"{ev['span_text']}\"  ({src} Page {ev['page_number']})",
                style="Intense Quote")
        if f.get("audit_trail"):
            doc.add_paragraph("Audit trail:", style="Heading 3")
            for stg in f["audit_trail"]:
                mark = "PASS" if stg["outcome"] == "pass" else "FAIL"
                doc.add_paragraph(f"[{mark}] {stg['stage']} — {stg['detail']}",
                                  style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
